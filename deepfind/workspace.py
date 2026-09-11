from __future__ import annotations

import base64
import json
import os
import queue
import re
import secrets
import subprocess
import threading
import time
from collections import deque
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from typing import Any, Iterator

from .coding_runtime import (
    CodingRuntimeError,
    RuntimeExecution,
    RuntimeInfo,
    _extract_workspace_archive,
)

WORKSPACE_ROOT = "/workspace"
MAX_TERMINALS = 4
MAX_TERMINAL_MESSAGE = 64 * 1024
MAX_SCROLLBACK_BYTES = 5 * 1024 * 1024
MAX_SCROLLBACK_EVENTS = 10_000
MAX_SUBSCRIBER_EVENTS = 256
MAX_DIRECTORY_ENTRIES = 500
MAX_TEXT_BYTES = 2 * 1024 * 1024
MAX_PDF_BYTES = 250 * 1024 * 1024
MAX_WORKBOOK_BYTES = 100 * 1024 * 1024
MAX_WORD_BYTES = 100 * 1024 * 1024
MAX_SHEET_CELLS = 10_000
PARSER_TIMEOUT = 30

_CHAT_ID_RE = re.compile(r"^chat_[a-zA-Z0-9_-]{1,64}$")
_TERMINAL_ID_RE = re.compile(r"^term_[a-f0-9]{32}$")
_BLOCKED_PARTS = frozenset(
    {
        ".git",
        ".env",
        ".ssh",
        ".aws",
        ".gnupg",
        ".config",
        ".docker",
        ".kube",
        "credentials",
        "secrets",
    }
)


class WorkspaceError(RuntimeError):
    def __init__(self, code: str, message: str, status_code: int = 400) -> None:
        super().__init__(message)
        self.code = code
        self.message = message
        self.status_code = status_code


@dataclass(frozen=True)
class WorkspaceConfig:
    enabled: bool
    runtime: str
    image: str
    seccomp_profile: Path


def validate_relative_path(value: str, *, allow_root: bool = True) -> str:
    if "\x00" in value or "\\" in value or "//" in value:
        raise WorkspaceError("path_not_allowed", "The requested path is not allowed", 403)
    if value in {"", "."}:
        if allow_root:
            return "."
        raise WorkspaceError("path_not_allowed", "The requested path is not allowed", 403)
    path = PurePosixPath(value)
    if path.is_absolute() or any(part in {"", ".", ".."} for part in path.parts):
        raise WorkspaceError("path_not_allowed", "The requested path is not allowed", 403)
    if any(part.startswith(".") or part.lower() in _BLOCKED_PARTS for part in path.parts):
        raise WorkspaceError("path_not_allowed", "The requested path is not allowed", 403)
    return path.as_posix()


def _container_name(chat_id: str) -> str:
    if not _CHAT_ID_RE.fullmatch(chat_id):
        raise WorkspaceError("workspace_not_found", "The chat workspace does not exist", 404)
    return f"deepfind-chat-{chat_id.removeprefix('chat_').lower()}"


_WORKSPACE_SCRIPT = r"""
import base64, csv, datetime, html, json, mimetypes, os, pathlib, stat, sys
root = pathlib.Path('/workspace')
mode, relative = sys.argv[1], sys.argv[2]
parts = [] if relative == '.' else relative.split('/')
if any(not p or p in ('.', '..') or p.startswith('.') for p in parts):
    raise PermissionError('path_not_allowed')
candidate = root.joinpath(*parts)
current = root
for part in parts:
    current = current / part
    info = os.lstat(current)
    if stat.S_ISLNK(info.st_mode):
        raise PermissionError('path_not_allowed')
resolved = candidate.resolve(strict=True)
resolved.relative_to(root)
info = os.stat(resolved, follow_symlinks=False)
def iso(ts):
    return datetime.datetime.fromtimestamp(ts, datetime.timezone.utc).isoformat()
def file_meta(path):
    item = os.stat(path, follow_symlinks=False)
    return {'name': path.name or 'workspace', 'path': '.' if path == root else path.relative_to(root).as_posix(),
            'size': item.st_size, 'modified_at': iso(item.st_mtime), 'is_directory': stat.S_ISDIR(item.st_mode),
            'mime_type': mimetypes.guess_type(path.name)[0] or 'application/octet-stream'}
if mode == 'metadata':
    print(json.dumps(file_meta(resolved), separators=(',', ':')))
elif mode == 'list':
    if not resolved.is_dir(): raise NotADirectoryError(relative)
    entries = []
    for path in sorted(resolved.iterdir(), key=lambda p: (not p.is_dir(), p.name.lower())):
        if path.name.startswith('.') or path.name.lower() in {'.git','.env','.ssh','.aws','.gnupg','.config','.docker','.kube','credentials','secrets'}:
            continue
        try:
            if path.is_symlink(): continue
            entries.append(file_meta(path))
        except OSError:
            continue
        if len(entries) >= 501: break
    print(json.dumps({'path': relative, 'entries': entries[:500], 'truncated': len(entries) > 500}, separators=(',', ':')))
elif mode == 'content':
    if not resolved.is_file(): raise FileNotFoundError(relative)
    limit = int(sys.argv[3])
    if info.st_size > limit: raise OverflowError('file_too_large')
    offset = int(sys.argv[4]) if len(sys.argv) > 4 else 0
    remaining = int(sys.argv[5]) if len(sys.argv) > 5 else info.st_size - offset
    if offset < 0 or remaining < 0 or offset > info.st_size: raise ValueError('invalid_range')
    fd = os.open(resolved, os.O_RDONLY | getattr(os, 'O_NOFOLLOW', 0))
    with os.fdopen(fd, 'rb') as source:
        source.seek(offset)
        while remaining:
            chunk = source.read(min(65536, remaining))
            if not chunk: break
            sys.stdout.buffer.write(chunk)
            remaining -= len(chunk)
elif mode == 'workbook':
    if resolved.suffix.lower() in ('.csv', '.tsv'):
        dialect = '\t' if resolved.suffix.lower() == '.tsv' else ','
        rows = columns = 0
        with resolved.open('r', encoding='utf-8-sig', errors='replace', newline='') as source:
            for row in csv.reader(source, delimiter=dialect):
                rows += 1; columns = max(columns, len(row))
        print(json.dumps({'file': file_meta(resolved), 'sheets': [{'id':'sheet_1','name':resolved.stem,'rows':rows,'columns':columns,'hidden':False,'freeze_panes':''}], 'warnings': []}, separators=(',', ':')))
        raise SystemExit
    from openpyxl import load_workbook
    if info.st_size > int(sys.argv[3]): raise OverflowError('file_too_large')
    book = load_workbook(resolved, read_only=True, data_only=False, keep_links=False)
    payload = {'file': file_meta(resolved), 'sheets': [], 'warnings': ['Formula results are stored values and may be stale.']}
    for index, sheet in enumerate(book.worksheets):
        payload['sheets'].append({'id': f'sheet_{index + 1}', 'name': sheet.title, 'rows': sheet.max_row,
                                  'columns': sheet.max_column, 'hidden': sheet.sheet_state != 'visible',
                                  'freeze_panes': str(getattr(sheet, 'freeze_panes', '') or '')})
    book.close()
    print(json.dumps(payload, separators=(',', ':')))
elif mode == 'sheet':
    from openpyxl.utils.cell import range_boundaries
    if info.st_size > int(sys.argv[3]): raise OverflowError('file_too_large')
    sheet_index, cell_range, max_cells = int(sys.argv[4]), sys.argv[5], int(sys.argv[6])
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    if (max_col-min_col+1)*(max_row-min_row+1) > max_cells: raise OverflowError('range_too_large')
    if resolved.suffix.lower() in ('.csv', '.tsv'):
        if sheet_index != 0: raise KeyError('sheet_not_found')
        dialect = '\t' if resolved.suffix.lower() == '.tsv' else ','
        cells = []
        with resolved.open('r', encoding='utf-8-sig', errors='replace', newline='') as source:
            rows = list(csv.reader(source, delimiter=dialect))
        for row_number in range(min_row, max_row + 1):
            source_row = rows[row_number - 1] if row_number <= len(rows) else []
            output = []
            for column in range(min_col, max_col + 1):
                value = source_row[column - 1] if column <= len(source_row) else ''
                output.append({'value': value, 'display': value, 'formula': None, 'number_format': ''})
            cells.append(output)
        print(json.dumps({'sheet_id':'sheet_1','range':cell_range,'cells':cells,'truncated':False}, separators=(',', ':')))
        raise SystemExit
    from openpyxl import load_workbook
    formulas = load_workbook(resolved, read_only=True, data_only=False, keep_links=False)
    values = load_workbook(resolved, read_only=True, data_only=True, keep_links=False)
    if sheet_index < 0 or sheet_index >= len(formulas.worksheets): raise KeyError('sheet_not_found')
    formula_sheet, value_sheet = formulas.worksheets[sheet_index], values.worksheets[sheet_index]
    cells = []
    for row in range(min_row, max_row + 1):
        output = []
        for column in range(min_col, max_col + 1):
            formula_cell = formula_sheet.cell(row, column)
            stored = value_sheet.cell(row, column).value
            formula = formula_cell.value if formula_cell.data_type == 'f' else None
            display = '' if stored is None else str(stored)
            output.append({'value': stored, 'display': display, 'formula': formula, 'number_format': formula_cell.number_format})
        cells.append(output)
    formulas.close(); values.close()
    print(json.dumps({'sheet_id': f'sheet_{sheet_index + 1}', 'range': cell_range, 'cells': cells, 'truncated': False}, separators=(',', ':'), default=str))
elif mode == 'word':
    from docx import Document
    if info.st_size > int(sys.argv[3]): raise OverflowError('file_too_large')
    document = Document(resolved)
    chunks, outline = [], []
    for index, paragraph in enumerate(document.paragraphs):
        text = html.escape(paragraph.text)
        if not text: continue
        style = (paragraph.style.name if paragraph.style else '').lower()
        if style.startswith('heading'):
            digits = ''.join(ch for ch in style if ch.isdigit())
            level = min(6, max(1, int(digits or '1')))
            heading_id = f'h{index}'
            outline.append({'id': heading_id, 'level': level, 'text': paragraph.text})
            chunks.append(f'<h{level} id="{heading_id}">{text}</h{level}>')
        else:
            chunks.append(f'<p>{text}</p>')
    for table in document.tables:
        rows = []
        for row in table.rows:
            rows.append('<tr>' + ''.join(f'<td>{html.escape(cell.text)}</td>' for cell in row.cells) + '</tr>')
        chunks.append('<table><tbody>' + ''.join(rows) + '</tbody></table>')
    title = document.core_properties.title or resolved.stem
    print(json.dumps({'title': title, 'outline': outline, 'html': ''.join(chunks), 'warnings': []}, separators=(',', ':')))
elif mode == 'manifest':
    files = {}
    for path in root.rglob('*'):
        try:
            relative_path = path.relative_to(root)
            if any(part.startswith('.') for part in relative_path.parts) or path.is_symlink() or not path.is_file():
                continue
            item = os.stat(path, follow_symlinks=False)
            files[relative_path.as_posix()] = [item.st_size, item.st_mtime_ns]
        except OSError:
            continue
        if len(files) >= 10000: raise OverflowError('workspace_too_large')
    print(json.dumps(files, separators=(',', ':')))
else:
    raise ValueError('unsupported operation')
"""


class TerminalSession:
    def __init__(self, terminal_id: str, chat_id: str, process: subprocess.Popen[bytes]) -> None:
        self.id = terminal_id
        self.chat_id = chat_id
        self.process = process
        self.created_at = time.time()
        self.sequence = 0
        self.exit_code: int | None = None
        self._events: deque[dict[str, Any]] = deque()
        self._event_bytes = 0
        self._subscribers: set[queue.Queue[dict[str, Any]]] = set()
        self._lock = threading.Lock()
        self._reader = threading.Thread(target=self._read_output, daemon=True)
        self._reader.start()

    def _publish(self, event: dict[str, Any]) -> None:
        encoded_size = len(json.dumps(event, separators=(",", ":")).encode("utf-8"))
        with self._lock:
            self.sequence += 1
            message = {**event, "sequence": self.sequence}
            self._events.append(message)
            self._event_bytes += encoded_size
            while len(self._events) > MAX_SCROLLBACK_EVENTS or self._event_bytes > MAX_SCROLLBACK_BYTES:
                removed = self._events.popleft()
                self._event_bytes -= len(json.dumps(removed, separators=(",", ":")).encode("utf-8"))
            subscribers = tuple(self._subscribers)
        for subscriber in subscribers:
            try:
                subscriber.put_nowait(message)
            except queue.Full:
                try:
                    subscriber.get_nowait()
                    subscriber.put_nowait(
                        {
                            "type": "error",
                            "code": "terminal_backpressure",
                            "message": "Terminal output paused because the client is too slow",
                            "sequence": self.sequence,
                        }
                    )
                except (queue.Empty, queue.Full):
                    pass
                with self._lock:
                    self._subscribers.discard(subscriber)

    def _read_output(self) -> None:
        assert self.process.stdout is not None
        for raw_line in iter(self.process.stdout.readline, b""):
            try:
                message = json.loads(raw_line.decode("utf-8"))
                if message.get("type") == "output":
                    message["data"] = base64.b64decode(message["data"]).decode("utf-8", errors="replace")
                if message.get("type") == "exit":
                    self.exit_code = message.get("exit_code")
                self._publish(message)
            except (ValueError, TypeError, KeyError):
                self._publish({"type": "error", "code": "terminal_protocol_error", "message": "Invalid terminal event"})
        return_code = self.process.wait()
        if self.exit_code is None:
            self.exit_code = return_code
            self._publish({"type": "exit", "exit_code": return_code, "signal": None})

    def send(self, message: dict[str, Any]) -> None:
        raw = json.dumps(message, separators=(",", ":")).encode("utf-8") + b"\n"
        if len(raw) > MAX_TERMINAL_MESSAGE:
            raise WorkspaceError("terminal_message_too_large", "The terminal message is too large")
        if self.process.poll() is not None or self.process.stdin is None:
            raise WorkspaceError("terminal_not_found", "Terminal session no longer exists", 404)
        if message.get("type") == "input":
            data = str(message.get("data", "")).encode("utf-8")
            message = {"type": "input", "data": base64.b64encode(data).decode("ascii")}
            raw = json.dumps(message, separators=(",", ":")).encode("utf-8") + b"\n"
        self.process.stdin.write(raw)
        self.process.stdin.flush()

    def subscribe(self, after: int = 0) -> tuple[queue.Queue[dict[str, Any]], list[dict[str, Any]]]:
        subscriber: queue.Queue[dict[str, Any]] = queue.Queue(maxsize=MAX_SUBSCRIBER_EVENTS)
        with self._lock:
            replay = [event for event in self._events if int(event["sequence"]) > after]
            self._subscribers.add(subscriber)
        return subscriber, replay

    def unsubscribe(self, subscriber: queue.Queue[dict[str, Any]]) -> None:
        with self._lock:
            self._subscribers.discard(subscriber)

    def close(self) -> None:
        if self.process.poll() is None:
            try:
                self.send({"type": "close"})
                self.process.wait(timeout=3)
            except (WorkspaceError, subprocess.TimeoutExpired):
                self.process.terminate()
                try:
                    self.process.wait(timeout=3)
                except subprocess.TimeoutExpired:
                    self.process.kill()

    def metadata(self) -> dict[str, Any]:
        return {
            "terminal_id": self.id,
            "chat_id": self.chat_id,
            "cwd": ".",
            "status": "exited" if self.process.poll() is not None else "ready",
            "exit_code": self.exit_code,
            "sequence": self.sequence,
        }


class ChatContainerManager:
    def __init__(self, config: WorkspaceConfig) -> None:
        self.config = config
        self._terminals: dict[str, TerminalSession] = {}
        self._coding_locks: dict[str, threading.Lock] = {}
        self._lock = threading.Lock()

    def info(self) -> RuntimeInfo:
        if not self.config.enabled:
            raise WorkspaceError("workspace_unavailable", "Chat workspaces are disabled", 503)
        result = self._run(
            [self.config.runtime, "version", "--format", "{{json .}}"],
            timeout=15,
            error_code="workspace_unavailable",
        )
        return RuntimeInfo(
            runtime=self.config.runtime,
            version=result.stdout.decode("utf-8", errors="replace").strip(),
            image=self.config.image,
            image_digest=self.config.image.rsplit("@", 1)[-1],
        )

    def ensure(self, chat_id: str) -> None:
        self.info()
        name = _container_name(chat_id)
        inspected = subprocess.run(
            [self.config.runtime, "inspect", name],
            capture_output=True,
            timeout=10,
            check=False,
        )
        if inspected.returncode == 0:
            state = json.loads(inspected.stdout or b"[]")[0]["State"]
            if not state.get("Running"):
                self._run([self.config.runtime, "start", name], timeout=20)
            return
        bridge = Path(__file__).with_name("workspace_terminal_bridge.py").resolve()
        runner = Path(__file__).with_name("coding_runner.py").resolve()
        command = [
            self.config.runtime,
            "create",
            "--name",
            name,
            "--label",
            f"deepfind.chat_id={chat_id}",
            "--network",
            "none",
            "--read-only",
            "--user",
            "65532:65532",
            "--cap-drop",
            "ALL",
            "--security-opt",
            "no-new-privileges",
            "--security-opt",
            f"seccomp={self.config.seccomp_profile.resolve()}",
            "--pids-limit",
            "128",
            "--cpus",
            "2",
            "--memory",
            "1g",
            "--memory-swap",
            "1g",
            "--tmpfs",
            "/workspace:rw,nosuid,nodev,size=512m,uid=65532,gid=65532,mode=0700",
            "--tmpfs",
            "/tmp:rw,nosuid,nodev,noexec,size=128m,uid=65532,gid=65532,mode=0700",
            "--mount",
            f"type=bind,source={bridge},target=/opt/deepfind/terminal_bridge.py,readonly",
            "--mount",
            f"type=bind,source={runner},target=/opt/deepfind/coding_runner.py,readonly",
            "--env",
            "HOME=/workspace",
            "--env",
            "PATH=/opt/venv-template/bin:/usr/local/bin:/usr/bin:/bin",
            "--env",
            "MPLCONFIGDIR=/workspace/.mplcache",
            "--entrypoint",
            "sleep",
            self.config.image,
            "infinity",
        ]
        self._run(command, timeout=30, error_code="workspace_unavailable")
        try:
            self._run([self.config.runtime, "start", name], timeout=20, error_code="workspace_unavailable")
        except Exception:
            subprocess.run([self.config.runtime, "rm", "-f", name], capture_output=True, check=False)
            raise

    def remove(self, chat_id: str) -> None:
        name = _container_name(chat_id)
        with self._lock:
            terminal_ids = [key for key, value in self._terminals.items() if value.chat_id == chat_id]
            terminals = [self._terminals.pop(key) for key in terminal_ids]
            self._coding_locks.pop(chat_id, None)
        for terminal in terminals:
            terminal.close()
        result = subprocess.run(
            [self.config.runtime, "rm", "-f", name],
            capture_output=True,
            timeout=20,
            check=False,
        )
        text = (result.stdout + result.stderr).decode("utf-8", errors="replace").lower()
        if result.returncode != 0 and "no such container" not in text:
            raise WorkspaceError("workspace_cleanup_failed", "The chat workspace could not be removed", 500)

    def status(self, chat_id: str) -> dict[str, Any]:
        name = _container_name(chat_id)
        result = subprocess.run(
            [self.config.runtime, "inspect", name],
            capture_output=True,
            timeout=10,
            check=False,
        )
        if result.returncode != 0:
            return {"available": False, "status": "missing", "chat_id": chat_id}
        data = json.loads(result.stdout or b"[]")[0]
        return {
            "available": bool(data["State"].get("Running")),
            "status": "running" if data["State"].get("Running") else "stopped",
            "chat_id": chat_id,
            "container_id": str(data.get("Id", ""))[:12],
            "workspace_root": ".",
        }

    def list_files(self, chat_id: str, relative: str = ".") -> dict[str, Any]:
        return self._json_operation(chat_id, "list", relative)

    def metadata(self, chat_id: str, relative: str) -> dict[str, Any]:
        return self._json_operation(chat_id, "metadata", relative, allow_root=False)

    def read_file(
        self,
        chat_id: str,
        relative: str,
        *,
        limit: int = MAX_TEXT_BYTES,
        offset: int = 0,
        length: int | None = None,
    ) -> bytes:
        normalized = validate_relative_path(relative, allow_root=False)
        args = [str(limit), str(offset)]
        if length is not None:
            args.append(str(length))
        result = self._operation(chat_id, "content", normalized, *args, timeout=PARSER_TIMEOUT)
        return result.stdout

    def workbook(self, chat_id: str, relative: str) -> dict[str, Any]:
        return self._json_operation(chat_id, "workbook", relative, str(MAX_WORKBOOK_BYTES), allow_root=False)

    def sheet(self, chat_id: str, relative: str, sheet_id: str, cell_range: str) -> dict[str, Any]:
        if not re.fullmatch(r"sheet_[1-9][0-9]*", sheet_id):
            raise WorkspaceError("document_parse_failed", "The worksheet does not exist", 404)
        if not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]*:[A-Z]{1,3}[1-9][0-9]*", cell_range.upper()):
            raise WorkspaceError("document_parse_failed", "The worksheet range is invalid")
        index = int(sheet_id.removeprefix("sheet_")) - 1
        return self._json_operation(
            chat_id,
            "sheet",
            relative,
            str(MAX_WORKBOOK_BYTES),
            str(index),
            cell_range.upper(),
            str(MAX_SHEET_CELLS),
            allow_root=False,
        )

    def word(self, chat_id: str, relative: str) -> dict[str, Any]:
        return self._json_operation(chat_id, "word", relative, str(MAX_WORD_BYTES), allow_root=False)

    def create_terminal(self, chat_id: str) -> dict[str, Any]:
        self.ensure(chat_id)
        with self._lock:
            active = sum(
                1 for terminal in self._terminals.values()
                if terminal.chat_id == chat_id and terminal.process.poll() is None
            )
            if active >= MAX_TERMINALS:
                raise WorkspaceError("terminal_limit_reached", "The terminal limit has been reached", 409)
        process = subprocess.Popen(
            [
                self.config.runtime,
                "exec",
                "-i",
                _container_name(chat_id),
                "python",
                "/opt/deepfind/terminal_bridge.py",
            ],
            stdin=subprocess.PIPE,
            stdout=subprocess.PIPE,
            stderr=subprocess.DEVNULL,
        )
        terminal = TerminalSession(f"term_{secrets.token_hex(16)}", chat_id, process)
        with self._lock:
            self._terminals[terminal.id] = terminal
        return terminal.metadata()

    def list_terminals(self, chat_id: str) -> list[dict[str, Any]]:
        with self._lock:
            return [item.metadata() for item in self._terminals.values() if item.chat_id == chat_id]

    def get_terminal(self, chat_id: str, terminal_id: str) -> TerminalSession:
        if not _TERMINAL_ID_RE.fullmatch(terminal_id):
            raise WorkspaceError("terminal_not_found", "Terminal session no longer exists", 404)
        with self._lock:
            terminal = self._terminals.get(terminal_id)
        if terminal is None or terminal.chat_id != chat_id:
            raise WorkspaceError("terminal_not_found", "Terminal session no longer exists", 404)
        return terminal

    def close_terminal(self, chat_id: str, terminal_id: str) -> None:
        terminal = self.get_terminal(chat_id, terminal_id)
        terminal.close()
        with self._lock:
            self._terminals.pop(terminal_id, None)

    def run_coding(self, chat_id: str, request: dict[str, Any], timeout: int) -> dict[str, Any]:
        self.ensure(chat_id)
        task_id = str(request["task_id"])
        request_path = f"/tmp/deepfind-request-{task_id}.json"
        payload = json.dumps(request, ensure_ascii=False).encode("utf-8")
        write_script = "import pathlib,sys; pathlib.Path(sys.argv[1]).write_bytes(sys.stdin.buffer.read())"
        written = subprocess.run(
            [
                self.config.runtime,
                "exec",
                "-i",
                _container_name(chat_id),
                "python",
                "-c",
                write_script,
                request_path,
            ],
            input=payload,
            capture_output=True,
            timeout=10,
            check=False,
        )
        if written.returncode != 0:
            raise CodingRuntimeError("sandbox_unavailable", "The coding request could not be prepared")
        result = subprocess.run(
            [
                self.config.runtime,
                "exec",
                "--env",
                f"DEEPFIND_REQUEST_PATH={request_path}",
                _container_name(chat_id),
                "python",
                "/opt/deepfind/coding_runner.py",
            ],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=timeout,
            check=False,
        )
        cleanup = (
            "import pathlib,sys;"
            "p=pathlib.Path(sys.argv[1]);"
            "p.unlink(missing_ok=True)"
        )
        subprocess.run(
            [self.config.runtime, "exec", _container_name(chat_id), "python", "-c", cleanup, request_path],
            capture_output=True,
            timeout=5,
            check=False,
        )
        result_path = "/workspace/.deepfind/result.json"
        loaded = subprocess.run(
            [self.config.runtime, "exec", _container_name(chat_id), "cat", result_path],
            capture_output=True,
            timeout=10,
            check=False,
        )
        if loaded.returncode != 0:
            raise CodingRuntimeError("invalid_result", "The coding result was not produced")
        return json.loads(loaded.stdout)

    def run_coding_task(
        self,
        chat_id: str,
        request: dict[str, Any],
        timeout: int,
    ) -> tuple[dict[str, Any], bytes | None]:
        with self._lock:
            coding_lock = self._coding_locks.setdefault(chat_id, threading.Lock())
        with coding_lock:
            result = self.run_coding(chat_id, request, timeout)
            artifacts = result.get("artifacts", [])
            if not isinstance(artifacts, list) or not all(isinstance(path, str) for path in artifacts):
                raise CodingRuntimeError("invalid_result", "The coding artifact manifest was invalid")
            archive = self.export_files(chat_id, artifacts) if artifacts else None
            return result, archive

    def manifest(self, chat_id: str) -> dict[str, list[int]]:
        return self._json_operation(chat_id, "manifest", ".")

    def export_files(self, chat_id: str, paths: list[str]) -> bytes:
        self.ensure(chat_id)
        normalized = [validate_relative_path(path, allow_root=False) for path in paths]
        encoded = base64.b64encode(json.dumps(normalized).encode("utf-8")).decode("ascii")
        command = (
            "import base64,json,pathlib,sys,tarfile;"
            "root=pathlib.Path('/workspace');"
            "paths=json.loads(base64.b64decode(sys.argv[1]));"
            "bundle=tarfile.open(fileobj=sys.stdout.buffer,mode='w|');"
            "\nfor relative in paths:\n"
            " p=root.joinpath(*relative.split('/'))\n"
            " if p.is_symlink() or not p.is_file(): raise ValueError('invalid artifact')\n"
            " p.resolve(strict=True).relative_to(root)\n"
            " bundle.add(p,arcname=relative,recursive=False)\n"
            "bundle.close()"
        )
        result = subprocess.run(
            [self.config.runtime, "exec", _container_name(chat_id), "python", "-c", command, encoded],
            capture_output=True,
            timeout=PARSER_TIMEOUT,
            check=False,
        )
        if result.returncode != 0:
            raise CodingRuntimeError("invalid_result", "The coding artifacts could not be exported")
        return result.stdout

    def _json_operation(
        self,
        chat_id: str,
        mode: str,
        relative: str,
        *args: str,
        allow_root: bool = True,
    ) -> dict[str, Any]:
        normalized = validate_relative_path(relative, allow_root=allow_root)
        result = self._operation(chat_id, mode, normalized, *args, timeout=PARSER_TIMEOUT)
        try:
            return json.loads(result.stdout)
        except json.JSONDecodeError as exc:
            raise WorkspaceError("document_parse_failed", "The document could not be parsed", 422) from exc

    def _operation(
        self,
        chat_id: str,
        mode: str,
        relative: str,
        *args: str,
        timeout: int,
    ) -> subprocess.CompletedProcess[bytes]:
        self.ensure(chat_id)
        try:
            result = subprocess.run(
                [
                    self.config.runtime,
                    "exec",
                    _container_name(chat_id),
                    "python",
                    "-c",
                    _WORKSPACE_SCRIPT,
                    mode,
                    relative,
                    *args,
                ],
                capture_output=True,
                timeout=timeout,
                check=False,
            )
        except subprocess.TimeoutExpired as exc:
            raise WorkspaceError("document_parse_failed", "The document parser timed out", 422) from exc
        if result.returncode == 0:
            return result
        error = result.stderr.decode("utf-8", errors="replace")
        if "PermissionError" in error:
            raise WorkspaceError("path_not_allowed", "The requested path is not allowed", 403)
        if "FileNotFoundError" in error or "NotADirectoryError" in error:
            raise WorkspaceError("file_not_found", "The requested file does not exist", 404)
        if "OverflowError" in error:
            code = "file_too_large" if "file_too_large" in error else "document_parse_failed"
            raise WorkspaceError(code, "The requested file or range exceeds the configured limit", 413)
        if "ModuleNotFoundError" in error or "ImportError" in error:
            raise WorkspaceError("unsupported_format", "The container image lacks the required document parser", 415)
        raise WorkspaceError("document_parse_failed", "The document could not be parsed", 422)

    def _run(
        self,
        command: list[str],
        *,
        timeout: int,
        error_code: str = "workspace_unavailable",
    ) -> subprocess.CompletedProcess[bytes]:
        try:
            result = subprocess.run(command, capture_output=True, timeout=timeout, check=False)
        except (FileNotFoundError, subprocess.TimeoutExpired) as exc:
            raise WorkspaceError(error_code, "The container workspace is unavailable", 503) from exc
        if result.returncode != 0:
            raise WorkspaceError(error_code, "The container workspace is unavailable", 503)
        return result


class ChatContainerCodingRuntime:
    def __init__(self, manager: ChatContainerManager, chat_id: str) -> None:
        self.manager = manager
        self.chat_id = chat_id

    def probe(self) -> RuntimeInfo:
        return self.manager.info()

    def run(
        self,
        workspace: Path,
        task_id: str,
        *,
        timeout: int | None = None,
    ) -> RuntimeExecution:
        request_path = workspace / ".deepfind" / "request.json"
        request = json.loads(request_path.read_text(encoding="utf-8"))
        result, archive = self.manager.run_coding_task(self.chat_id, request, timeout or 120)
        if archive is not None:
            _extract_workspace_archive(archive, workspace)
        control = workspace / ".deepfind"
        control.mkdir(exist_ok=True)
        (control / "result.json").write_text(
            json.dumps(result, ensure_ascii=False, separators=(",", ":")),
            encoding="utf-8",
        )
        return RuntimeExecution(
            exit_code=0 if result.get("ok") else 1,
            stdout="",
            stderr="",
        )


def stream_terminal_events(
    terminal: TerminalSession,
    *,
    after: int = 0,
) -> Iterator[dict[str, Any]]:
    subscriber, replay = terminal.subscribe(after)
    try:
        yield from replay
        while terminal.process.poll() is None or not subscriber.empty():
            try:
                yield subscriber.get(timeout=0.5)
            except queue.Empty:
                continue
    finally:
        terminal.unsubscribe(subscriber)
